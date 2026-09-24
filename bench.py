# -*- coding: utf-8 -*-
"""
全流程测试台：随机抽 6 张真实样本，统计耗时 / 行数 / 数值规整数，
并输出前若干行内容便于肉眼核对。
用法: python bench.py [样本数] [--dir 目录]
"""
import sys, os, time, random, glob
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import core

SRC_DIR = r"C:\Users\toxuj\Desktop\FAPIAO"
N = 6
if len(sys.argv) > 1 and sys.argv[1].isdigit():
    N = int(sys.argv[1])
if "--dir" in sys.argv:
    SRC_DIR = sys.argv[sys.argv.index("--dir") + 1]

files = sorted(glob.glob(os.path.join(SRC_DIR, "*")))
files = [f for f in files if os.path.splitext(f)[1].lower() in core.IMG_EXTS]
random.seed(20260924)
picked = random.sample(files, min(N, len(files)))

print(f"样本目录: {SRC_DIR}")
print(f"候选 {len(files)} 个，抽取 {len(picked)} 个\n")

total_t = 0.0
summary = []
for i, f in enumerate(picked, 1):
    name = os.path.basename(f)
    t0 = time.time()
    rows, errors = core.convert_many([f])
    dt = time.time() - t0
    total_t += dt
    nums = sum(1 for r in rows for c in r[2:] if isinstance(c, (int, float)))
    cols = max((len(r) for r in rows), default=0)
    summary.append((name, len(rows), cols, nums, dt, len(errors)))
    print(f"[{i}/{len(picked)}] {name[:44]}")
    print(f"      {len(rows)} 行 x {max(cols-2,0)} 列 | 规整数值 {nums} 个 | {dt:.1f}s | 异常 {len(errors)}")

print("\n" + "=" * 72)
print(f"合计 {len(picked)} 个文件，总耗时 {total_t:.1f}s，平均 {total_t/max(len(picked),1):.1f}s/文件")
print(f"总行数 {sum(s[1] for s in summary)}，总规整数值 {sum(s[3] for s in summary)}，总异常 {sum(s[5] for s in summary)}")

out = r"D:\workbuddy\杂项\img22excel\bench_out.xlsx"
all_rows, all_err = core.convert_many(picked)
if all_rows:
    core.write_excel(all_rows, out, picked, keep_raw=True)
    print(f"\n汇总输出: {out}")

    import openpyxl
    wb = openpyxl.load_workbook(out); ws = wb.active
    print(f"Excel 规模: {ws.max_row} 行 x {ws.max_column} 列")
    print("\n前 6 行预览（前 7 列）:")
    for r in range(1, min(ws.max_row, 7) + 1):
        vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column, 7) + 1)]
        print("  ", [str(v)[:20] if v is not None else None for v in vals])
